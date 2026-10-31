"""Treasurer screen (owner in Claude): expenses with receipt photo, taxes, report per period, per truck, taxes to remit, Excel export."""
import json, os, datetime, re, sys, base64, io
from playwright.sync_api import sync_playwright
import openpyxl
SC = os.path.dirname(os.path.abspath(__file__))
PAGE = open('/home/claude/garage/carnet.html', encoding='utf8').read()
MOCK = open(os.path.join(SC, 'mock.js'), encoding='utf8').read()
XLSX = open(os.path.join(SC, 'xlsxlib/node_modules/xlsx/dist/xlsx.full.min.js'), encoding='utf8').read()
RESET = 'html{color-scheme:light}:root{padding-top:env(safe-area-inset-top,0px);padding-bottom:env(safe-area-inset-bottom,0px)}body{margin:0;font:14px system-ui}img{max-width:100%}[hidden]{display:none!important}'
T = datetime.date.today(); iso = lambda d: d.isoformat(); MK = iso(T)[:7]
MONTHS = ['Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin', 'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']
MLABEL = MONTHS[T.month - 1] + ' ' + str(T.year)
seed = {'config/main': {'biz': {'name': 'Garage NextLine'}, 'mechanics': ['Saleh', 'Ali'], 'rate': 120, 'taxesOn': True, 'prefix': 'F-', 'startNo': 1,
                        'trucks': [{'id': 'c1', 'name': 'Camion 1', 'on': True, 'mech': 'Saleh'}, {'id': 'c2', 'name': 'Camion 2', 'on': False, 'mech': ''}, {'id': 'c3', 'name': 'Camion 3', 'on': False, 'mech': ''}]}}
def inv(no, sub, paid):
    tps, tvq = round(sub * 0.05, 2), round(sub * 0.09975, 2)
    return {'number': no, 'sub': sub, 'tps': tps, 'tvq': tvq, 'total': round(sub + tps + tvq, 2), 'paid': paid, 'method': 'Débit' if paid else '', 'issuedAt': iso(T), 'lines': [{'type': 'mo', 'desc': 'Travail', 'qty': 1, 'price': sub}], 'taxesOn': True}
seed['jobs/r1'] = {'date': iso(T), 'time': '09:00', 'mechanic': 'Saleh', 'truck': 'c1', 'status': 'termine', 'client': {'name': 'Marie Tremblay'}, 'services': ['Freins'], 'invoice': inv('F-0001', 1000.0, True)}
seed['jobs/r2'] = {'date': iso(T), 'time': '13:00', 'mechanic': 'Saleh', 'truck': 'c1', 'status': 'termine', 'client': {'name': 'Paul Roy'}, 'services': ['Huile'], 'invoice': inv('F-0002', 500.0, False)}
seed['jobs/x1'] = {'date': iso(T), 'time': '15:00', 'mechanic': 'Saleh', 'truck': 'c1', 'status': 'annule', 'client': {'name': 'Zoé Gagnon'}, 'services': ['Pneus'], 'lost': {'amount': -200, 'reason': 'client', 'note': 'reporté'}}
seed['jobs/n1'] = {'date': iso(T), 'time': '16:00', 'mechanic': 'Saleh', 'truck': 'c1', 'status': 'noshow', 'client': {'name': 'Luc Bélanger'}, 'services': ['Diagnostic']}
def harness(mode, ui='full'):
    pre = 'try{localStorage.setItem("gm.ui",%s)}catch(e){}' % json.dumps(ui) if ui else ''
    return ('<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">'
            f'<style>{RESET}</style><script>{pre}window.__MOCK={json.dumps(mode)};</script><script>{MOCK}</script></head><body>{PAGE}</body></html>')
results = []
def check(name, cond, detail=''):
    results.append((name, bool(cond))); print(('PASS ' if cond else 'FAIL ') + name + (('  -> ' + str(detail)[:300]) if detail and not cond else ''), flush=True)
def open_page(br, mode, width=1100, ui='full', scheme='light'):
    ctx = br.new_context(viewport={'width': width, 'height': 900}, color_scheme=scheme, timezone_id='America/Toronto', locale='fr-CA', has_touch=width < 700, **({'user_agent': 'Mozilla/5.0 (iPhone; CPU iPhone OS 17_5 like Mac OS X) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.5 Mobile/15E148 Safari/604.1'} if width < 700 else {}))
    pg = ctx.new_page(); errs = []
    pg.on('pageerror', lambda e: errs.append(str(e)))
    pg.on('console', lambda m: errs.append(m.text) if m.type == 'error' and 'fonts.g' not in m.text and 'ERR_FAILED' not in m.text else None)
    def route(r):
        u = r.request.url
        if u == 'https://local.test/': r.fulfill(status=200, content_type='text/html; charset=utf-8', body=harness(mode, ui))
        elif u.endswith('/xlsx/0.18.5/xlsx.full.min.js'): r.fulfill(status=200, content_type='application/javascript', body=XLSX)
        else: r.abort()
    pg.route('**/*', route); pg.goto('https://local.test/'); pg.wait_for_timeout(700)
    return ctx, pg, errs
def txt(pg, sel): return (pg.locator(sel).first.text_content() or '').replace('\xa0', ' ').replace(' ', ' ')
def store(pg, path): return pg.evaluate('p => window.__store[p] || null', path)
def exps(pg): return pg.evaluate("Object.entries(window.__store).filter(([k]) => k.startsWith('expenses/')).map(([k, v]) => Object.assign({id: k.slice(9)}, v))")
def wait(pg, ms=300): pg.wait_for_timeout(ms)
def no_hscroll(pg): return pg.evaluate('document.documentElement.scrollWidth <= document.documentElement.clientWidth + 1')
def m(x): return x.replace('\xa0', ' ').replace(' ', ' ')

with sync_playwright() as p:
    br = p.chromium.launch()
    ctx, pg, errs = open_page(br, {'seed': seed})
    tabs = pg.evaluate("[...document.querySelectorAll('#tabbar .tab:not([hidden])')].map(b => b.textContent.trim())")
    check('owner: « Trésorerie » tab next to Réglages', tabs == ['Accueil', 'Agenda', 'Clients', 'Factures', 'Trésorerie', 'Réglages'], tabs)
    pg.click('.tab[data-t="tresor"]'); wait(pg, 500)
    check('Trésorerie opens on Dépenses, this month', txt(pg, '#tr-label') == MLABEL and pg.locator('[data-act="tr-view"][data-v="dep"].on').count() == 1, txt(pg, '#tr-label'))
    check('no expense yet: empty card + « + Ajouter une dépense »', pg.locator('.empty-card [data-act="exp-new"]').count() == 1)
    # 1st expense: NAPA 114,98 taxes included, Camion 1, debit, receipt photo
    pg.click('.top-in [data-act="exp-new"]'); wait(pg)
    check('new expense form: date today, taxes included by default', pg.input_value('#x-date') == iso(T) and pg.is_checked('#x-taxin'))
    pg.fill('#x-total', '114,98'); wait(pg, 100)
    check('taxes computed from the total (TPS 5,00 · TVQ 9,98 · avant taxes 100,00 $)', pg.input_value('#x-tps') == '5,00' and pg.input_value('#x-tvq') == '9,98' and m(txt(pg, '#x-sub')) == '100,00 $', [pg.input_value('#x-tps'), pg.input_value('#x-tvq'), txt(pg, '#x-sub')])
    pg.fill('#x-vendor', 'NAPA'); pg.select_option('#x-cat', 'pieces'); pg.select_option('#x-truck', 'c1'); pg.click('[data-act="x-pay"][data-v="Débit"]'); pg.fill('#x-desc', 'Plaquettes Civic')
    pg.set_input_files('#x-file', os.path.join(SC, 'recu_test.jpg')); wait(pg, 1500)
    check('receipt photo made smaller, preview shown', pg.locator('#x-prev img.rcthumb').count() == 1 and 'Photo prête' in txt(pg, '#x-prev'), txt(pg, '#x-prev'))
    pg.click('#x-save'); wait(pg, 700)
    e1 = [e for e in exps(pg) if e.get('vendor') == 'NAPA']
    e1 = e1[0] if e1 else {}
    check('saved: 114,98 = 100,00 + 5,00 + 9,98, Camion 1, Débit, note', e1.get('total') == 114.98 and e1.get('sub') == 100.0 and e1.get('tps') == 5.0 and e1.get('tvq') == 9.98 and e1.get('truck') == 'c1' and e1.get('method') == 'Débit' and e1.get('desc') == 'Plaquettes Civic' and e1.get('cat') == 'pieces', e1)
    rc = store(pg, e1.get('receipt') or '-') or {}
    check('receipt stored apart (JPEG, under 330 Ko of text)', e1.get('receipt') == 'receipts/' + e1.get('id', '') and str(rc.get('img', '')).startswith('data:image/jpeg') and len(rc.get('img', '')) < 330000, len(rc.get('img', '')))
    check('sheet closed, list shows NAPA with « Reçu ✓ »', not pg.locator('#sheet').is_visible() and 'NAPA' in txt(pg, '#exp-list') and 'Reçu ✓' in txt(pg, '#exp-list'))
    # 2nd: Esso 80 $ taxes included, no receipt · 3rd: insurance 300 $ no taxes
    pg.click('.top-in [data-act="exp-new"]'); wait(pg); pg.fill('#x-total', '80'); pg.fill('#x-vendor', 'Esso'); pg.select_option('#x-cat', 'essence'); pg.click('#x-save'); wait(pg, 600)
    pg.click('.top-in [data-act="exp-new"]'); wait(pg); pg.fill('#x-total', '300'); pg.uncheck('#x-taxin'); wait(pg, 100)
    check('no taxes: TPS/TVQ hidden, avant taxes = total', pg.locator('#x-taxes').is_hidden())
    pg.fill('#x-vendor', 'Intact Assurance'); pg.select_option('#x-cat', 'assur'); pg.click('#x-save'); wait(pg, 600)
    es = {e['vendor']: e for e in exps(pg)}
    check('Esso: 80 = 69,58 + 3,48 + 6,94 · insurance: 300 without taxes', es['Esso']['sub'] == 69.58 and es['Esso']['tps'] == 3.48 and es['Esso']['tvq'] == 6.94 and es['Intact Assurance']['tps'] == 0 and es['Intact Assurance']['sub'] == 300, (es['Esso'], es['Intact Assurance']))
    check('list: 3 expenses, total paid, 2 without receipt', '3 dépenses · 494,98 $ payés' in m(txt(pg, '#tr-count')) and '2 sans reçu' in txt(pg, '#tr-count'), m(txt(pg, '#tr-count')))
    pg.click('.top-in [data-act="exp-new"]'); wait(pg); pg.fill('#x-total', '10'); pg.fill('#x-tps', '6'); pg.fill('#x-tvq', '5'); pg.click('#x-save'); wait(pg, 300)
    check('taxes bigger than the total refused', 'plus petites que le total' in txt(pg, '#x-err') and len(exps(pg)) == 3)
    pg.click('.sheet-head [data-act="sheet-close"]'); wait(pg)
    # edit NAPA: total changes → taxes follow; reopening keeps saved taxes
    pg.click('.xrow[data-id="%s"]' % e1['id']); wait(pg)
    check('reopening keeps the saved values', pg.input_value('#x-total') == '114,98' and pg.input_value('#x-tps') == '5,00' and pg.input_value('#x-vendor') == 'NAPA' and pg.locator('[data-act="x-pay"].on').text_content() == 'Débit')
    pg.fill('#x-total', '229,95'); wait(pg, 100); pg.click('#x-save'); wait(pg, 600)
    e1b = [e for e in exps(pg) if e['id'] == e1['id']][0]
    check('edited: 229,95 = 200,00 + 10,00 + 19,95, receipt kept', e1b['total'] == 229.95 and e1b['sub'] == 200.0 and e1b['tps'] == 10.0 and e1b['tvq'] == 19.95 and e1b['receipt'] == e1['receipt'], e1b)
    pg.click('.xrow[data-id="%s"]' % e1['id']); wait(pg); pg.click('[data-act="exp-receipt"]'); wait(pg, 500)
    check('« Voir le reçu » shows the photo', pg.locator('#rc-img').count() == 1 and pg.get_attribute('#rc-img', 'src').startswith('data:image/jpeg'))
    pg.click('#sheet-body [data-act="exp-open"]'); wait(pg)
    check('« Retour à la dépense »', pg.locator('#expform').count() == 1)
    pg.click('.sheet-head [data-act="sheet-close"]'); wait(pg)
    # report
    pg.click('[data-act="tr-view"][data-v="rap"]'); wait(pg, 300)
    res = m(txt(pg, '#rep-res'))
    # revenue 1500 ; expenses 200 + 69.58 + 300 = 569.58 ; profit 930.42
    check('report: revenue 1 500,00 $, expenses 569,58 $, profit 930,42 $', '1 500,00 $' in res and '569,58 $' in res and '930,42 $' in res, res)
    check('summary tiles agree', '930,42 $' in m(txt(pg, '#tr-sum')))
    cat = pg.evaluate("[...document.querySelectorAll('#rep-cat .hbar')].map(h => h.querySelector('span').textContent + '=' + h.querySelector('b').textContent.replace(/\\s/g,' '))")
    check('expenses by category, biggest first', cat == ['Assurances=300,00 $', 'Pièces=200,00 $', 'Essence=69,58 $'], cat)
    trk = pg.evaluate("[...document.querySelectorAll('#rep-truck tbody tr')].map(r => [...r.children].map(c => c.textContent.replace(/\\s/g,' ')))")
    check('per truck: Camion 1 revenue 1 500, expenses 200 (parts), profit 1 300', trk[0] == ['Camion 1', '1 500,00 $', '200,00 $', '1 300,00 $'], trk)
    tax = pg.evaluate("[...document.querySelectorAll('#rep-tax tbody tr')].map(r => [...r.children].map(c => c.textContent.replace(/\\s/g,' ')))")
    # TPS: collected 50 + 25 = 75 ; paid 10 + 3.48 = 13.48 ; due 61.52 · TVQ: 99.75 + 49.88 = 149.63 ; paid 19.95 + 6.94 = 26.89 ; due 122.74
    check('taxes to remit: TPS 75,00 − 13,48 = 61,52 ; TVQ 149,63 − 26,89 = 122,74', tax[0][1:] == ['75,00 $', '149,63 $'] and tax[1][1:] == ['−13,48 $', '−26,89 $'] and tax[2][1:] == ['61,52 $', '122,74 $'], tax)
    check('cancellations and no-shows, receivables', '-200,00 $' in m(txt(pg, '#rep-other')) and '1 no-show' in txt(pg, '#rep-other') and '574,88 $' in m(txt(pg, '#rep-other')), m(txt(pg, '#rep-other')))
    pg.click('[data-act="tr-xlsx"]'); wait(pg, 2500)
    saved = pg.evaluate('window.__saved')
    check('Excel file saved for the accountant', len(saved) == 1 and saved[0]['filename'] == 'Trésorerie %s.xlsx' % MLABEL, [s['filename'] for s in saved])
    if saved:
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(saved[0]['b64'])), data_only=True)
        check('Excel: 4 sheets', wb.sheetnames == ['Résumé', 'Revenus', 'Dépenses', 'Annulations'], wb.sheetnames)
        rs = {r[0].value: [c.value for c in r[1:]] for r in wb['Résumé'].iter_rows() if r[0].value}
        check('Excel summary: revenue, expenses, profit, taxes', rs['Revenus (factures émises)'][:4] == [1500, 75, 149.63, 1724.63] and rs['Dépenses'][0] == 569.58 and rs['Profit (avant impôts)'][0] == 930.42 and rs['TPS à remettre'][0] == 61.52 and rs['TVQ à remettre'][0] == 122.74, rs)
        dp = [[c.value for c in r] for r in wb['Dépenses'].iter_rows()]
        check('Excel expenses: 3 lines with category, truck, receipt', len(dp) == 4 and dp[0][:3] == ['Date', 'Fournisseur', 'Catégorie'] and any(r[1] == 'NAPA' and r[3] == 'Camion 1' and r[10] == 'oui' and r[8] == 229.95 for r in dp[1:]), dp)
        rv = [[c.value for c in r] for r in wb['Revenus'].iter_rows()]
        check('Excel revenue: 2 invoices', len(rv) == 3 and {r[1] for r in rv[1:]} == {'F-0001', 'F-0002'}, rv)
        an = [[c.value for c in r] for r in wb['Annulations'].iter_rows()]
        check('Excel cancellations: -200, reason', len(an) == 2 and an[1][2] == -200 and an[1][3] == 'Le client a annulé', an)
    # periods
    pg.click('[data-act="tr-kind"][data-v="quarter"]'); wait(pg)
    q = (T.month - 1) // 3 + 1
    check('quarter label', txt(pg, '#tr-label') == ('1er' if q == 1 else '%de' % q) + ' trimestre %d' % T.year, txt(pg, '#tr-label'))
    pg.click('[data-act="tr-kind"][data-v="year"]'); wait(pg)
    check('year label, same totals', txt(pg, '#tr-label') == 'Année %d' % T.year and '930,42 $' in m(txt(pg, '#tr-sum')))
    pg.click('[data-act="tr-shift"][data-v="-1"]'); wait(pg)
    check('previous year: nothing', txt(pg, '#tr-label') == 'Année %d' % (T.year - 1) and '0,00 $' in m(txt(pg, '#tr-sum')))
    pg.click('[data-act="tr-kind"][data-v="month"]'); wait(pg)
    check('back to a month (January of that year)', txt(pg, '#tr-label') == 'Janvier %d' % (T.year - 1), txt(pg, '#tr-label'))
    pg.evaluate("() => { const S = window.__garageApi.S; S.trP = null; S.trView = 'dep'; window.__garageApi.render(); }"); wait(pg)
    # delete (asks first), receipt goes with it — back to the original state
    ins = [e for e in exps(pg) if e['vendor'] == 'Intact Assurance'][0]
    pg.click('.xrow[data-id="%s"]' % ins['id']); wait(pg); pg.click('[data-act="exp-del"]'); wait(pg, 200)
    check('delete asks first', pg.locator('[data-act="exp-del-no"]').count() == 1 and len(exps(pg)) == 3)
    pg.click('[data-act="exp-del-no"]'); wait(pg, 100); pg.click('[data-act="exp-del"]'); wait(pg, 100); pg.click('[data-act="exp-del"]'); wait(pg, 500)
    check('deleted', len(exps(pg)) == 2 and not pg.locator('#sheet').is_visible())
    pg.click('.xrow[data-id="%s"]' % e1['id']); wait(pg); pg.click('[data-act="exp-del"]'); wait(pg, 100); pg.click('[data-act="exp-del"]'); wait(pg, 500)
    check('deleting NAPA removes its receipt photo too', store(pg, e1['receipt']) is None and len(exps(pg)) == 1)
    esso = exps(pg)[0]; pg.click('.xrow[data-id="%s"]' % esso['id']); wait(pg); pg.click('[data-act="exp-del"]'); wait(pg, 100); pg.click('[data-act="exp-del"]'); wait(pg, 500)
    check('all test expenses removed (original state)', exps(pg) == [] and pg.evaluate("Object.keys(window.__store).filter(k => k.startsWith('receipts/')).length") == 0)
    check('no page errors', not errs, errs)
    ctx.close()

    # ---------- phone ----------
    ctx, pg, errs = open_page(br, {'seed': seed}, width=390, ui=None)
    tabs = pg.evaluate("[...document.querySelectorAll('#tabbar .tab:not([hidden])')].map(b => b.dataset.t || '+')")
    check('phone (owner): tab bar unchanged', sorted(tabs) == sorted(['jour', 'clients', '+', 'factures', 'plus']), tabs)
    pg.click('.tab[data-t="plus"]'); wait(pg)
    check('phone Plus: « Trésorerie »', pg.locator('nav.ph-menu a[href="#tresor"]').count() == 1)
    pg.click('.tab-new'); wait(pg)
    check('phone « + »: « Nouvelle dépense »', pg.locator('[data-act="ph-go"][data-v="exp"]').count() == 1)
    pg.click('[data-act="ph-go"][data-v="exp"]'); wait(pg)
    check('phone: expense form opens, camera allowed for the receipt', pg.locator('#expform').count() == 1 and pg.get_attribute('#x-file', 'capture') == 'environment')
    pg.fill('#x-total', '45'); pg.fill('#x-vendor', 'Canadian Tire'); pg.click('#x-save'); wait(pg, 600)
    check('phone: expense saved', len(exps(pg)) == 1)
    pg.click('.tab[data-t="plus"]'); wait(pg); pg.click('nav.ph-menu a[href="#tresor"]'); wait(pg, 500)
    check('phone: Trésorerie, no sideways scroll, back arrow to Plus', pg.locator('#exp-list .xrow').count() == 1 and no_hscroll(pg) and pg.locator('.tab[data-t="plus"][aria-current="page"]').count() == 1, no_hscroll(pg))
    pg.click('[data-act="tr-view"][data-v="rap"]'); wait(pg)
    check('phone: report fits the screen', no_hscroll(pg))
    pg.screenshot(path=os.path.join(SC, 'p_tresor_rapport.png'), full_page=True)
    pg.evaluate("() => { for (const k of Object.keys(window.__store)) if (k.startsWith('expenses/')) delete window.__store[k]; }")
    check('no page errors (phone)', not errs, errs)
    ctx.close()
    # ---------- not the owner ----------
    ctx, pg, errs = open_page(br, {'seed': seed, 'notAdmin': True})
    check('a contributor who is not the owner: no Trésorerie', pg.locator('.tab[data-t="tresor"]').is_hidden())
    ctx.close()
    br.close()
ok = sum(1 for _, c in results if c)
print('%d/%d passed' % (ok, len(results)))
sys.exit(0 if ok == len(results) else 1)
